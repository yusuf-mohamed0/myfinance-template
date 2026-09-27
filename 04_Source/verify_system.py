# -*- coding: utf-8 -*-
r"""
VERIFY_SYSTEM - end-to-end check that every part of the MyFinance template
is present, wired and unmodified. Prints an ASCII checklist (PASS/FAIL) and
exits non-zero on any failure.

Checks:
  1. Folder structure (6 numbered folders)
  2. config.json present and valid (03_System/config.json)
  3. Key files exist (SMS input, scripts, reports, site, docs)
  4. Scripts compile (py_compile)
  5. SYSTEM LOCK: every 04_Source/*.py + init_project.py matches the
     SHA-256 recorded in 03_System/source_manifest.json.
     Any edit by anyone other than the system owner = FAIL
     "SYSTEM FILES MODIFIED" (works even without git).
  6. Data chain: sms_raw*.txt -> summary_<year>.json -> plan_YYYY-MM.xlsx
     (data checks SKIP when no summary yet / txns == 0)
  7. summary numbers self-consistent
  8. Website: lock gate + gate hash matches config + market pulse +
     income+expense quick-log with inline edit + zero emoji/arrow chars
  9. market_watch.json complete + weekly refresh cadence
 10. ACL lockdown (Windows only - icacls; skipped on macOS/Linux)

Usage: python verify_system.py
ASCII-only console output.
"""
import glob
import hashlib
import json
import os
import py_compile
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import mfconfig
from mfconfig import CFG, GATE, YEAR

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
fails = []
passes = 0
skips = 0


def check(name, ok, detail=""):
    global passes
    if ok:
        passes += 1
        print("[PASS] " + name)
    else:
        fails.append(name)
        print("[FAIL] " + name + ("  -> " + detail if detail else ""))


def skip(name, why):
    global skips
    skips += 1
    print("[SKIP] " + name + "  (" + why + ")")


# ---------------------------------------------------------------- 1. folders
for d in ["01_Data", "02_Reports", "03_System", "04_Source", "05_Docs", "06_Web"]:
    check("folder " + d, os.path.isdir(os.path.join(BASE, d)))

# ---------------------------------------------------------------- 2. config
cfg_ok = os.path.isfile(os.path.join(BASE, "03_System", "config.json"))
check("config.json exists", cfg_ok, "run init_project.py first")
check("config passcode is not the placeholder",
      str(CFG.get("passcode", "")) not in ("", "change-me"))
check("config has a display name", bool(CFG.get("name")))

# ---------------------------------------------------------------- 3. files
MONTH = None
plan_files = sorted(glob.glob(os.path.join(BASE, "02_Reports", "plan_*.xlsx")))
if plan_files:
    MONTH = os.path.basename(plan_files[-1])[len("plan_"):-len(".xlsx")]

FILES = [
    ("02_Reports", "market_watch.json"),
    ("03_System", "My_Financial_System.xlsx"),
    ("04_Source", "analyze_sms.py"),
    ("04_Source", "build_excel.py"),
    ("04_Source", "build_word.py"),
    ("04_Source", "plan_month.py"),
    ("04_Source", "daily_status.py"),
    ("04_Source", "build_web.py"),
    ("04_Source", "publish_web.py"),
    ("04_Source", "verify_system.py"),
    ("04_Source", "mfconfig.py"),
    ("05_Docs", "README.md"),
    ("init_project.py", None),
]
for folder, fn in FILES:
    if fn is None:
        p = os.path.join(BASE, folder)
    else:
        p = os.path.join(BASE, folder, fn)
    ok = os.path.isfile(p) and os.path.getsize(p) > 0
    label = folder if fn is None else folder + "/" + fn
    check("file " + label, ok, "missing or empty")

sms_inputs = glob.glob(os.path.join(BASE, "01_Data", "sms_raw*.txt"))
check("sms input present (01_Data/sms_raw*.txt)", bool(sms_inputs),
      "paste your bank SMS (see README 'Add your SMS')")

# ---------------------------------------------------------------- 4. compile
SCRIPTS = ["analyze_sms.py", "build_excel.py", "build_word.py",
           "plan_month.py", "daily_status.py", "build_web.py",
           "publish_web.py", "verify_system.py", "mfconfig.py"]
for sc in SCRIPTS:
    p = os.path.join(BASE, "04_Source", sc)
    try:
        py_compile.compile(p, doraise=True)
        check("compile " + sc, True)
    except Exception as e:
        check("compile " + sc, False, str(e))

# ---------------------------------------------------------------- 5. SYSTEM LOCK
# sha256 of every shipped script; regenerate ONLY by the system owner
# (yusuf) after an approved change - see README "System lock policy".
mf_path = os.path.join(BASE, "03_System", "source_manifest.json")
if not os.path.isfile(mf_path):
    check("source_manifest.json exists", False,
          "system files cannot be authenticated - reinstall template")
else:
    try:
        manifest = json.load(open(mf_path, encoding="utf-8"))
        expected = manifest.get("files", manifest)
        tampered = []
        missing = []
        for rel, want in sorted(expected.items()):
            p = os.path.join(BASE, rel.replace("/", os.sep))
            if not os.path.isfile(p):
                missing.append(rel)
                continue
            # normalize line endings so CRLF checkouts (Windows git) are
            # not flagged as tampering - content, not EOL, is what we lock
            raw = open(p, "rb").read().replace(b"\r\n", b"\n")
            got = hashlib.sha256(raw).hexdigest()
            if got != want:
                tampered.append(rel)
        check("SYSTEM FILES UNMODIFIED (manifest)", not tampered and not missing,
              "SYSTEM FILES MODIFIED: " + ", ".join(tampered + missing))
    except Exception as e:
        check("source_manifest.json reads", False, str(e))

# ---------------------------------------------------------------- 6. data chain
summary_path = os.path.join(BASE, "02_Reports", "summary_{}.json".format(YEAR))
s = None
if not os.path.isfile(summary_path):
    skip("summary_{}.json".format(YEAR), "run analyze_sms.py first")
else:
    try:
        s = json.load(open(summary_path, encoding="utf-8"))
        check("summary_{}.json parses".format(YEAR), True)
    except Exception as e:
        check("summary_{}.json parses".format(YEAR), False, str(e))

# ---------------------------------------------------------------- 7. numbers
if s is not None:
    if (s.get("n_txns") or 0) == 0:
        skip("summary numeric checks", "txns == 0 (no SMS parsed yet)")
    else:
        deb, ref, net = s.get("total_debits"), s.get("total_refunds"), s.get("net_out")
        check("net_out = debits - refunds",
              abs((deb or 0) - (ref or 0) - (net or 0)) < 0.01,
              "{} - {} != {}".format(deb, ref, net))
        check("txns > 0", (s.get("n_txns") or 0) > 0, str(s.get("n_txns")))
        check("period present", bool(s.get("period")), "")
        cats = s.get("categories") or {}
        check("categories mapped >= 3", len(cats) >= 3, str(len(cats)))
        m = s.get("monthly") or {}
        md = round(sum(v.get("debit", 0) for v in m.values()), 2)
        check("monthly debits sum == total_debits", abs(md - (deb or 0)) < 0.02,
              "{} vs {}".format(md, deb))

if plan_files:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(plan_files[-1])
        names = wb.sheetnames
        check("plan has >=3 sheets", len(names) >= 3, str(names))
        ws = wb[wb.sheetnames[0]]
        check("plan first cell has title", bool(ws.cell(1, 1).value))
    except Exception as e:
        check("plan xlsx readable", False, str(e))
else:
    skip("plan xlsx", "run plan_month.py <income> first")

# filled workbook copy (only exists after analyze_sms ran)
filled = os.path.join(BASE, "03_System",
                      "My_Financial_System_{}_SMS.xlsx".format(YEAR))
if os.path.isfile(filled):
    try:
        import openpyxl
        wb2 = openpyxl.load_workbook(filled)
        check("filled system workbook 8 sheets", len(wb2.sheetnames) == 8,
              str(wb2.sheetnames))
    except Exception as e:
        check("filled system workbook readable", False, str(e))
else:
    skip("filled system workbook", "run analyze_sms.py first")

# ---------------------------------------------------------------- 8. website
idx = os.path.join(BASE, "06_Web", "index.html")
if not os.path.isfile(idx):
    skip("website", "run build_web.py first")
else:
    try:
        h = open(idx, encoding="utf-8").read()
        check("web index has lock gate", 'id="lock"' in h)
        check("web gate hash matches config passcode", GATE in h,
              "expected sha256 of config.json passcode")
        check("web has plan table", h.count("<tr>") >= 14, str(h.count("<tr>")))
        tok_files = [f for f in os.listdir(os.path.join(BASE, "06_Web"))
                     if len(f) > 11 and f[10] == "_" and f[:10].isalnum()]
        check("web has >=5 token-named downloads", len(tok_files) >= 5,
              str(len(tok_files)))
        check("web has market pulse section",
              'id="market"' in h and h.count('class="mk-i"') >= 9,
              str(h.count('class="mk-i"')))
        check("web supports income + expense logging",
              'id="seg_in"' in h and 'id="seg_out"' in h and 'st_in' in h)
        check("web supports inline edit of entries",
              all(x in h for x in ("function editEntry", "function saveEdit",
                                   "function cancelEdit", "function setEdType")))
        emoji = sum(1 for c in h if (0x1F000 <= ord(c) <= 0x1FAFF)
                    or (0x2600 <= ord(c) <= 0x27BF)
                    or ord(c) in (0xFE0F, 0x2192, 0x2190))
        check("web has zero emoji/arrow chars", emoji == 0, str(emoji))
    except Exception as e:
        check("web index readable", False, str(e))

# ---------------------------------------------------------------- 9. market
mw_path = os.path.join(BASE, "02_Reports", "market_watch.json")
try:
    mw = json.load(open(mw_path, encoding="utf-8"))
    check("market_watch: >=6 items, impact, news, last_checked",
          len(mw.get("items", [])) >= 6 and len(mw.get("impact", [])) >= 3
          and len(mw.get("news", [])) >= 1 and bool(mw.get("last_checked")),
          "items={} impact={} news={}".format(
              len(mw.get("items", [])), len(mw.get("impact", [])),
              len(mw.get("news", []))))
    check("market refresh cadence weekly (json)", mw.get("refresh_cadence") == "weekly",
          str(mw.get("refresh_cadence")))
except Exception as e:
    check("market_watch.json parses", False, str(e))

if os.path.isfile(idx):
    try:
        hh = open(idx, encoding="utf-8").read()
        check("market weekly note on site", "أسبوعيًا" in hh)
        pm_src = open(os.path.join(BASE, "04_Source", "plan_month.py"),
                      encoding="utf-8").read()
        check("plan_month links market prices (market_link)",
              "market_watch" in pm_src and "market_link" in pm_src)
    except Exception as e:
        check("plan-market link check ran", False, str(e))

# ---------------------------------------------------------------- 10. ACL
if os.name == "nt":
    try:
        import subprocess
        out = subprocess.run(["icacls", BASE], capture_output=True, text=True,
                             encoding="utf-8", errors="replace").stdout
        lines = [l.strip() for l in out.splitlines()
                 if "(F)" in l or "(M)" in l or "(RX)" in l]
        # Windows defaults (SYSTEM, BUILTIN\Administrators) are OK; wide
        # grants (Everyone / Users / Authenticated Users) are the real leak.
        check("ACL: no Everyone/Users/Authenticated rights",
              bool(lines) and not any(
                  any(bad in l for bad in ("Everyone", "BUILTIN\\Users",
                                           "Authenticated Users"))
                  for l in lines),
              repr(lines[:6]))
    except Exception as e:
        check("ACL check ran", False, str(e))
else:
    skip("ACL lockdown", "non-Windows platform (set folder perms yourself)")

# ---------------------------------------------------------------- result
print("-" * 60)
print("RESULT: {} passed, {} failed, {} skipped".format(passes, len(fails), skips))
if fails:
    for f in fails:
        print("  broken: " + f)
    sys.exit(1)
print("ALL GREEN - system fully wired")
