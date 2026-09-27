# -*- coding: utf-8 -*-
r"""
BUILD_WEB - generates a self-contained static site in 06_Web from
summary_<year>.json + plan_YYYY-MM.xlsx, then copies the downloadable
reports next to it. Publish with publish_web.py.

Design: professional RTL Arabic dashboard, inline Lucide icon sprite
(no CDN, works offline), responsive mobile-first layout, no emojis.
Quick-log supports expenses (withdrawals) AND income entries,
stored in browser localStorage, with inline edit of any entry.

Usage: python build_web.py [plan_month YYYY-MM]   (default: current month)
Console output is ASCII-only.
Settings (passcode, name, city, year) come from 03_System/config.json
via mfconfig.py - nothing personal is hardcoded here.
"""
import json, os, shutil, sys, datetime as dt
from string import Template

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import mfconfig
from mfconfig import CFG, GATE, TOKEN_SALT, YEAR

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WEB = os.path.join(BASE, "06_Web")
SUMMARY = os.path.join(BASE, "02_Reports", "summary_{}.json".format(YEAR))

MONTH = sys.argv[1] if len(sys.argv) > 1 else dt.date.today().strftime("%Y-%m")
PLAN = os.path.join(BASE, "02_Reports", "plan_{}.xlsx".format(MONTH))

DL = [  # report files copied into the site for download
    ("dashboard_{}.html".format(YEAR), "dashboard_{}.html".format(YEAR), "لوحة التحكم التفاعلية"),
    ("Financial_Analysis_{}.xlsx".format(YEAR), "Financial_Analysis_{}.xlsx".format(YEAR), "تقرير التحليل الكامل"),
    ("transactions_{}.csv".format(YEAR), "transactions_{}.csv".format(YEAR), "كل المعاملات (CSV)"),
    ("summary_{}.json".format(YEAR), "summary_{}.json".format(YEAR), "الأرقام المجمعة (JSON)"),
    ("plan_{}.xlsx".format(MONTH), "plan_{}.xlsx".format(MONTH), "خطة الشهر (إكسل)"),
    ("market_watch.json", "market_watch.json", "أسعار السوق والأخبار (JSON)"),
]

# ---------------------------------------------------------------- Lucide icons
# Inner SVG of official Lucide icons (ISC license), embedded for offline use.
ICONS = {
"lock": '''<rect width="18" height="11" x="3" y="11" rx="2" ry="2"/><path d="M7 11V7a5 5 0 0 1 10 0v4"/>''',
"shield-check": '''<path d="M20 13c0 5-3.5 7.5-7.66 8.95a1 1 0 0 1-.67-.01C7.5 20.5 4 18 4 13V6a1 1 0 0 1 1-1c2 0 4.5-1.2 6.24-2.72a1.17 1.17 0 0 1 1.52 0C14.51 3.81 17 5 19 5a1 1 0 0 1 1 1z"/><path d="m9 12 2 2 4-4"/>''',
"wallet": '''<path d="M19 7V4a1 1 0 0 0-1-1H5a2 2 0 0 0 0 4h15a1 1 0 0 1 1 1v4h-3a2 2 0 0 0 0 4h3a1 1 0 0 1-1 1v3a1 1 0 0 1-1 1H5a2 2 0 0 1-2-2V5"/><path d="M3 5v14a2 2 0 0 0 2 2h15a1 1 0 0 0 1-1v-4"/>''',
"trending-up": '''<path d="M16 7h6v6"/><path d="m22 7-8.5 8.5-5-5L2 17"/>''',
"trending-down": '''<path d="M16 17h6v-6"/><path d="m22 17-8.5-8.5-5 5L2 7"/>''',
"piggy-bank": '''<path d="M11 17h3v2a1 1 0 0 0 1 1h2a1 1 0 0 0 1-1v-3a3.16 3.16 0 0 0 2-2h1a1 1 0 0 0 1-1v-2a1 1 0 0 0-1-1h-1a5 5 0 0 0-2-4V3a4 4 0 0 0-3.2 1.6l-.3.4H11a6 6 0 0 0-6 6v1a5 5 0 0 0 2 4v3a1 1 0 0 0 1 1h2a1 1 0 0 0 1-1z"/><path d="M16 10h.01"/><path d="M2 8v1a2 2 0 0 0 2 2h1"/>''',
"gauge": '''<path d="m12 14 4-4"/><path d="M3.34 19a10 10 0 1 1 17.32 0"/>''',
"calendar-days": '''<path d="M8 2v3"/><path d="M16 2v3"/><rect x="3" y="3" width="18" height="18" rx="2"/><path d="M3 9h18"/><path d="M8 13h.01"/><path d="M12 13h.01"/><path d="M16 13h.01"/><path d="M8 17h.01"/><path d="M12 17h.01"/><path d="M16 17h.01"/>''',
"calendar-check": '''<path d="M8 2v3"/><path d="M16 2v3"/><rect x="3" y="3" width="18" height="18" rx="2"/><path d="M3 9h18"/><path d="m9 15 2 2 4-4"/>''',
"clipboard-list": '''<rect width="8" height="4" x="8" y="2" rx="1" ry="1"/><path d="M16 4h2a2 2 0 0 1 2 2v14a2 2 0 0 1-2 2H6a2 2 0 0 1-2-2V6a2 2 0 0 1 2-2h2"/><path d="M12 11h4"/><path d="M12 16h4"/><path d="M8 11h.01"/><path d="M8 16h.01"/>''',
"plus": '''<path d="M5 12h14"/><path d="M12 5v14"/>''',
"copy": '''<rect width="14" height="14" x="8" y="8" rx="2" ry="2"/><path d="M4 16c-1.1 0-2-.9-2-2V4c0-1.1.9-2 2-2h10c1.1 0 2 .9 2 2"/>''',
"download": '''<path d="M12 15V3"/><path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4"/><path d="m7 10 5 5 5-5"/>''',
"info": '''<circle cx="12" cy="12" r="10"/><path d="M12 16v-4"/><path d="M12 8h.01"/>''',
"target": '''<circle cx="12" cy="12" r="10"/><circle cx="12" cy="12" r="6"/><circle cx="12" cy="12" r="2"/>''',
"circle-check": '''<circle cx="12" cy="12" r="10"/><path d="m16 9-5.5 5.5L8 12"/>''',
"circle-alert": '''<circle cx="12" cy="12" r="10"/><line x1="12" x2="12" y1="8" y2="12"/><line x1="12" x2="12.01" y1="16" y2="16"/>''',
"triangle-alert": '''<path d="m21.73 18-8-14a2 2 0 0 0-3.48 0l-8 14A2 2 0 0 0 4 21h16a2 2 0 0 0 1.73-3"/><path d="M12 9v4"/><path d="M12 17h.01"/>''',
"banknote": '''<rect width="20" height="12" x="2" y="6" rx="2"/><circle cx="12" cy="12" r="2"/><path d="M6 12h.01M18 12h.01"/>''',
"fuel": '''<path d="M14 13h2a2 2 0 0 1 2 2v2a2 2 0 0 0 4 0v-6.998a2 2 0 0 0-.59-1.42L18 5"/><path d="M14 21V5a2 2 0 0 0-2-2H5a2 2 0 0 0-2 2v16"/><path d="M2 21h13"/><path d="M3 9h11"/>''',
"smartphone": '''<rect width="14" height="20" x="5" y="2" rx="2" ry="2"/><path d="M12 18h.01"/>''',
"shopping-cart": '''<path d="m2.05 2.05 1.099-.028a1 1 0 0 1 1.008.815l2.69 14.347A1 1 0 0 0 7.83 18H18"/><path d="M4.563 5h16.435a1 1 0 0 1 .981 1.204l-1.026 6.226A2 2 0 0 1 18.962 14H6.25"/><circle cx="18" cy="20" r="2"/><circle cx="8" cy="20" r="2"/>''',
"utensils": '''<path d="M3 2v7c0 1.1.9 2 2 2h4a2 2 0 0 0 2-2V2"/><path d="M7 2v20"/><path d="M21 15V2a5 5 0 0 0-5 5v6c0 1.1.9 2 2 2h3Zm0 0v7"/>''',
"coffee": '''<path d="M10 2v2"/><path d="M14 2v2"/><path d="M16 8a1 1 0 0 1 1 1v8a4 4 0 0 1-4 4H7a4 4 0 0 1-4-4V9a1 1 0 0 1 1-1h14a4 4 0 1 1 0 8h-1"/><path d="M6 2v2"/>''',
"shopping-basket": '''<path d="m15 11-1 9"/><path d="m19 11-4-7"/><path d="M2 11h20"/><path d="m3.5 11 1.6 7.4a2 2 0 0 0 2 1.6h9.8a2 2 0 0 0 2-1.6l1.7-7.4"/><path d="M4.5 15.5h15"/><path d="m5 11 4-7"/><path d="m9 11 1 9"/>''',
"heart-pulse": '''<path d="M2 9.5a5.5 5.5 0 0 1 9.591-3.676.56.56 0 0 0 .818 0A5.49 5.49 0 0 1 22 9.5c0 2.29-1.5 4-3 5.5l-5.492 5.313a2 2 0 0 1-3 .019L5 15c-1.5-1.5-3-3.2-3-5.5"/><path d="M3.22 13H9.5l.5-1 2 4.5 2-7 1.5 3.5h5.27"/>''',
"shirt": '''<path d="M20.38 3.46 16 2a4 4 0 0 1-8 0L3.62 3.46a2 2 0 0 0-1.34 2.23l.58 3.47a1 1 0 0 0 .99.84H6v10c0 1.1.9 2 2 2h8a2 2 0 0 0 2-2V10h2.15a1 1 0 0 0 .99-.84l.58-3.47a2 2 0 0 0-1.34-2.23z"/>''',
"refresh-cw": '''<path d="M3 12a9 9 0 0 1 9-9 9.75 9.75 0 0 1 6.74 2.74L21 8"/><path d="M21 3v5h-5"/><path d="M21 12a9 9 0 0 1-9 9 9.75 9.75 0 0 1-6.74-2.74L3 16"/><path d="M8 16H3v5"/>''',
"send": '''<path d="M14.536 21.686a.5.5 0 0 0 .937-.024l6.5-19a.496.496 0 0 0-.635-.635l-19 6.5a.5.5 0 0 0-.024.937l7.93 3.18a2 2 0 0 1 1.112 1.11z"/><path d="m21.854 2.147-10.94 10.939"/>''',
"file-spreadsheet": '''<path d="M6 22a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2h8a2.4 2.4 0 0 1 1.704.706l3.588 3.588A2.4 2.4 0 0 1 20 8v12a2 2 0 0 1-2 2z"/><path d="M14 2v5a1 1 0 0 0 1 1h5"/><path d="M8 13h2"/><path d="M14 13h2"/><path d="M8 17h2"/><path d="M14 17h2"/>''',
"file-text": '''<path d="M6 22a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2h8a2.4 2.4 0 0 1 1.704.706l3.588 3.588A2.4 2.4 0 0 1 20 8v12a2 2 0 0 1-2 2z"/><path d="M14 2v5a1 1 0 0 0 1 1h5"/><path d="M10 9H8"/><path d="M16 13H8"/><path d="M16 17H8"/>''',
"braces": '''<path d="M8 3H7a2 2 0 0 0-2 2v5a2 2 0 0 1-2 2 2 2 0 0 1 2 2v5c0 1.1.9 2 2 2h1"/><path d="M16 21h1a2 2 0 0 0 2-2v-5c0-1.1.9-2 2-2a2 2 0 0 1-2-2V5a2 2 0 0 0-2-2h-1"/>''',
"table": '''<path d="M12 3v18"/><rect width="18" height="18" x="3" y="3" rx="2"/><path d="M3 9h18"/><path d="M3 15h18"/>''',
"layout-dashboard": '''<rect width="7" height="9" x="3" y="3" rx="1"/><rect width="7" height="5" x="14" y="3" rx="1"/><rect width="7" height="9" x="14" y="12" rx="1"/><rect width="7" height="5" x="3" y="16" rx="1"/>''',
"list-checks": '''<path d="M13 5h8"/><path d="M13 12h8"/><path d="M13 19h8"/><path d="m3 17 2 2 4-4"/><path d="m3 7 2 2 4-4"/>''',
"square-pen": '''<path d="M12 3H5a2 2 0 0 0-2 2v14a2 2 0 0 0 2 2h14a2 2 0 0 0 2-2v-7"/><path d="M18.375 2.625a1 1 0 0 1 3 3l-9.013 9.014a2 2 0 0 1-.853.505l-2.873.84a.5.5 0 0 1-.62-.62l.84-2.873a2 2 0 0 1 .506-.852z"/>''',
"x": '''<path d="M18 6 6 18"/><path d="m6 6 12 12"/>''',
"log-in": '''<path d="m10 17 5-5-5-5"/><path d="M15 12H3"/><path d="M15 3h4a2 2 0 0 1 2 2v14a2 2 0 0 1-2 2h-4"/>''',
"circle-dot": '''<circle cx="12" cy="12" r="1"/><circle cx="12" cy="12" r="10"/>''',
"chevron-left": '''<path d="m15 18-6-6 6-6"/>''',
}

ICON_RULES = [  # label keyword -> icon (first match wins)
    ("ادخار", "piggy-bank"), ("فائض", "piggy-bank"), ("سحب", "banknote"),
    ("موبايل", "smartphone"), ("بنزين", "fuel"), ("مشتريات", "shopping-cart"),
    ("مطاعم", "utensils"), ("كافيهات", "coffee"), ("بقالة", "shopping-basket"),
    ("صحة", "heart-pulse"), ("تسوق", "shirt"), ("اشتراكات", "refresh-cw"),
    ("تحويلات", "send"), ("طوارئ", "triangle-alert"), ("غير مصنف", "circle-alert"),
]


def svg(name, cls="ic"):
    return ('<svg class="{c}" aria-hidden="true"><use href="#ic-{n}"/></svg>'
            .format(c=cls, n=name))


def icon_for(label):
    s = str(label)
    for kw, ic in ICON_RULES:
        if kw in s:
            return ic
    return "circle-dot"


def esc(s):
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def token_for(fn):
    """unguessable prefix for download links (sha256 of secret+name, 10 chars)"""
    import hashlib
    return hashlib.sha256((TOKEN_SALT + fn).encode("utf-8")).hexdigest()[:10]


def money(x):
    try:
        return "{:,.0f}".format(float(x))
    except (TypeError, ValueError):
        return esc(x)


def human(b):
    if b >= 1048576:
        return "{:.1f} MB".format(b / 1048576.0)
    return "{:.0f} KB".format(b / 1024.0)


def sprite():
    parts = ['<symbol id="ic-{n}" viewBox="0 0 24 24" fill="none" '
             'stroke="currentColor" stroke-width="2" stroke-linecap="round" '
             'stroke-linejoin="round">{b}</symbol>'.format(n=n, b=b)
             for n, b in ICONS.items()]
    return ('<svg xmlns="http://www.w3.org/2000/svg" style="display:none" '
            'aria-hidden="true">' + "".join(parts) + "</svg>")


def expand_icons(text):
    out = text
    for n in ICONS:
        out = out.replace("$i-" + n, svg(n))
    return out


TEMPLATE = r"""<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="theme-color" content="#0E2A47">
<title>$name — لوحة الثروة</title>
<style>
  :root{
    --navy:#0E2A47; --navy2:#163A5F; --gold:#C9A227;
    --green:#15803D; --green-bg:#ECFDF3;
    --red:#B91C1C; --red-bg:#FEF2F2;
    --blue:#1D4ED8; --blue-bg:#EFF4FF;
    --amber:#B7791F; --amber-bg:#FFFBEB;
    --bg:#F4F6F9; --surface:#fff; --line:#E4E9F0;
    --text:#16212E; --muted:#5D6B80;
    --radius:14px;
    --shadow:0 1px 2px rgba(14,42,71,.05), 0 10px 26px -16px rgba(14,42,71,.25);
  }
  *{box-sizing:border-box; margin:0; padding:0;}
  html{-webkit-text-size-adjust:100%;}
  body{
    font-family:"Segoe UI","Noto Sans Arabic","Helvetica Neue",Tahoma,Arial,sans-serif;
    background:var(--bg); color:var(--text); line-height:1.65; direction:rtl;
    -webkit-tap-highlight-color:transparent;
  }
  .ic{width:18px; height:18px; flex:none;}
  .num{font-variant-numeric:tabular-nums;}

  /* ---------- lock gate ---------- */
  #lock{
    position:fixed; inset:0; z-index:60; display:flex; align-items:center;
    justify-content:center; padding:20px;
    background:radial-gradient(1100px 520px at 50% -12%, #163A5F 0%, var(--navy) 52%, #071527 100%);
  }
  .lock-card{
    width:min(380px,100%); text-align:center; color:#fff;
    background:rgba(255,255,255,.06); border:1px solid rgba(255,255,255,.14);
    border-radius:20px; padding:34px 26px; backdrop-filter:blur(8px);
  }
  .lock-logo{
    width:58px; height:58px; margin:0 auto 14px; border-radius:17px;
    display:grid; place-items:center;
    background:rgba(201,162,39,.14); border:1px solid rgba(201,162,39,.42);
    color:var(--gold);
  }
  .lock-logo .ic{width:26px; height:26px;}
  .lock-card h1{font-size:1.25rem; font-weight:700;}
  .lock-card p{font-size:.87rem; color:rgba(255,255,255,.72); margin-top:5px;}
  #lock input{
    width:100%; margin-top:20px; padding:13px; border-radius:12px;
    border:1px solid rgba(255,255,255,.24); background:rgba(255,255,255,.08);
    color:#fff; text-align:center; font-size:1.05rem; letter-spacing:4px;
  }
  #lock input:focus{outline:none; border-color:var(--gold);}
  .lock-btn{
    margin-top:12px; width:100%; justify-content:center;
    background:var(--gold); color:#1A2333; border:0;
  }
  .lock-btn:hover{background:#D9B23A;}
  .lock-err{min-height:1.3em; font-size:.85rem; color:#FCA5A5; margin-top:10px;}
  .lock-hint{
    display:flex; gap:7px; align-items:center; justify-content:center;
    font-size:.75rem; color:rgba(255,255,255,.55); margin-top:16px;
  }

  /* ---------- header ---------- */
  header{background:linear-gradient(135deg, var(--navy) 0%, var(--navy2) 100%); color:#fff;}
  .hd{
    max-width:1040px; margin:0 auto; padding:18px 16px;
    display:flex; align-items:center; justify-content:space-between;
    gap:14px; flex-wrap:wrap;
  }
  .brand{display:flex; align-items:center; gap:12px;}
  .logo{
    width:44px; height:44px; border-radius:13px; display:grid; place-items:center;
    background:rgba(255,255,255,.11); border:1px solid rgba(255,255,255,.18);
    color:var(--gold);
  }
  .logo .ic{width:23px; height:23px;}
  .brand h1{font-size:1.22rem; font-weight:700; letter-spacing:-.2px;}
  .brand .sub{font-size:.8rem; color:rgba(255,255,255,.72);}
  .chip{
    display:inline-flex; align-items:center; gap:7px; font-size:.78rem;
    background:rgba(255,255,255,.1); border:1px solid rgba(255,255,255,.16);
    padding:7px 13px; border-radius:999px; color:rgba(255,255,255,.9);
  }
  .chip .ic{width:14px; height:14px; color:var(--gold);}

  /* ---------- layout ---------- */
  .wrap{max-width:1040px; margin:0 auto; padding:20px 16px 6px;}
  section{
    background:var(--surface); border:1px solid var(--line);
    border-radius:var(--radius); padding:18px; margin-bottom:16px;
    box-shadow:var(--shadow);
    animation:fade .35s ease both;
  }
  @keyframes fade{from{opacity:0; transform:translateY(6px);} to{opacity:1; transform:none;}}
  @media (prefers-reduced-motion:reduce){*{animation:none !important; transition:none !important;}}
  .sec-h{display:flex; align-items:center; gap:10px; margin-bottom:14px;}
  .sec-h .ic{width:19px; height:19px; color:var(--navy);}
  .sec-h h2{font-size:1.02rem; font-weight:700; color:var(--navy);}
  .sec-h .hint{margin-inline-start:auto; font-size:.76rem; color:var(--muted);}

  /* ---------- KPI cards ---------- */
  .cards{display:grid; grid-template-columns:repeat(auto-fit,minmax(210px,1fr));
         gap:14px; margin-bottom:16px;}
  .card{
    background:var(--surface); border:1px solid var(--line);
    border-radius:var(--radius); padding:16px; box-shadow:var(--shadow);
    animation:fade .35s ease both;
  }
  .ci{width:36px; height:36px; border-radius:11px; display:grid; place-items:center; margin-bottom:10px;}
  .ci .ic{width:19px; height:19px;}
  .tone-red{background:var(--red-bg); color:var(--red);}
  .tone-green{background:var(--green-bg); color:var(--green);}
  .tone-navy{background:#EAF0F7; color:var(--navy);}
  .tone-gold{background:var(--amber-bg); color:var(--amber);}
  .tone-blue{background:var(--blue-bg); color:var(--blue);}
  .card .v{font-size:1.5rem; font-weight:700; letter-spacing:-.4px;}
  .card .l{font-size:.82rem; color:var(--muted); margin-top:2px;}

  /* ---------- market pulse (prices & news) ---------- */
  .mk{display:grid; grid-template-columns:repeat(auto-fit,minmax(175px,1fr)); gap:10px;}
  .mk-i{
    background:#F8FAFD; border:1px solid var(--line); border-radius:12px;
    padding:12px 13px; display:flex; flex-direction:column; gap:3px; min-width:0;
  }
  .mk-l{display:flex; align-items:center; gap:6px; font-size:.76rem; color:var(--muted);}
  .mk-l .ic{width:14px; height:14px; color:var(--navy);}
  .mk-v{font-size:1.02rem; font-weight:700; color:var(--navy); font-variant-numeric:tabular-nums;}
  .mk-s{font-size:.7rem; color:#8A97A8;}
  .mk-cols{display:grid; grid-template-columns:1fr 1fr; gap:14px; margin-top:14px;}
  .mk-h{display:flex; align-items:center; gap:7px; font-size:.85rem;
         font-weight:700; color:var(--navy); margin-bottom:8px;}
  .mk-h .ic{width:15px; height:15px; color:var(--gold);}
  .mk-ul{list-style:none; display:flex; flex-direction:column; gap:8px;}
  .mk-ul li{
    display:flex; gap:8px; align-items:flex-start; font-size:.82rem; color:#3A4A5E;
    background:#F8FAFD; border:1px solid var(--line); border-radius:10px; padding:9px 11px;
  }
  .mk-ul li .ic{width:14px; height:14px; color:var(--muted); margin-top:3px; flex:none;}
  .mk-date{display:block; font-size:.72rem; color:#8A97A8; margin-top:2px;}

  /* ---------- quick log ---------- */
  .seg{
    display:inline-flex; gap:4px; padding:4px; margin-bottom:12px;
    background:#EDF1F7; border:1px solid var(--line); border-radius:12px;
    width:100%;
  }
  .seg-btn{
    flex:1; display:flex; align-items:center; justify-content:center; gap:8px;
    min-height:42px; padding:9px 14px; border:0; border-radius:9px;
    background:transparent; color:var(--muted); font-size:.92rem; font-weight:700;
    font-family:inherit; cursor:pointer;
  }
  .seg-btn.on{background:#fff; color:var(--navy); box-shadow:0 1px 3px rgba(14,42,71,.16);}
  .seg-btn.on.t-in{color:var(--green);}
  .form{display:flex; flex-wrap:wrap; gap:10px;}
  .form input{
    padding:11px 12px; min-height:44px; border:1px solid var(--line);
    border-radius:11px; font-size:.95rem; font-family:inherit; background:#fff;
    color:var(--text);
  }
  .form input:focus{outline:none; border-color:var(--blue);
    box-shadow:0 0 0 3px rgba(29,78,216,.14);}
  #lg_d{width:165px;}
  #lg_a{width:150px;}
  #lg_n{flex:1 1 180px; min-width:0;}
  .btn{
    display:inline-flex; align-items:center; justify-content:center; gap:8px;
    min-height:44px; padding:11px 18px; border-radius:11px; border:1px solid transparent;
    font-size:.93rem; font-weight:700; font-family:inherit; cursor:pointer;
  }
  .btn-pri{background:var(--navy); color:#fff;}
  .btn-pri:hover{background:var(--navy2);}
  .btn-out{background:#fff; border-color:var(--line); color:var(--muted);}
  .btn-out:hover{border-color:#C9D3E0; color:var(--text);}
  .btn-danger{background:#fff; border-color:#F0C9C9; color:var(--red);}
  .btn-danger:hover{background:var(--red-bg);}
  .stats{display:grid; grid-template-columns:repeat(auto-fit,minmax(165px,1fr));
         gap:10px; margin-top:14px;}
  .stat{background:#F8FAFD; border:1px solid var(--line); border-radius:12px; padding:12px 13px;}
  .stat .sl{display:flex; align-items:center; gap:6px; font-size:.76rem; color:var(--muted);}
  .stat .sl .ic{width:14px; height:14px;}
  .stat .sv{font-size:1.16rem; font-weight:700; margin-top:3px;}
  .bar{height:7px; background:#E6ECF3; border-radius:99px; margin-top:9px; overflow:hidden;}
  .bar i{display:block; height:100%; border-radius:99px; width:0;
         background:var(--green); transition:width .35s ease;}
  .bar i.warn{background:var(--gold);}
  .bar i.over{background:var(--red);}
  .v-green{color:var(--green);} .v-red{color:var(--red);}
  .log-list{margin-top:6px;}
  .log-row{
    display:flex; align-items:center; gap:11px; padding:11px 2px;
    border-bottom:1px solid var(--line);
  }
  .log-row:last-child{border-bottom:0;}
  .lic{width:34px; height:34px; border-radius:10px; display:grid; place-items:center; flex:none;}
  .lic .ic{width:17px; height:17px;}
  .li-main{flex:1; min-width:0;}
  .li-t{font-size:.92rem; font-weight:600;}
  .li-s{font-size:.77rem; color:var(--muted);}
  .li-amt{font-weight:700; font-variant-numeric:tabular-nums; white-space:nowrap; font-size:.95rem;}
  .del{
    background:none; border:0; color:#9AA7B8; cursor:pointer;
    padding:8px; border-radius:9px; display:grid; place-items:center;
  }
  .del:hover{background:var(--red-bg); color:var(--red);}
  .del .ic{width:16px; height:16px;}
  .del.ed:hover{background:var(--blue-bg); color:var(--blue);}
  .del.ok{color:var(--green);}
  .del.ok:hover{background:var(--green-bg); color:var(--green);}

  /* ---------- inline edit row ---------- */
  .log-row.editing{background:#F8FAFD; align-items:flex-start; border-radius:10px; padding:12px 8px;}
  .log-row.editing .lic{margin-top:2px;}
  .ed-f{display:flex; gap:6px; flex-wrap:wrap; margin-bottom:7px;}
  .ed-f input{
    flex:1 1 84px; min-width:0; padding:7px 9px; font:inherit; font-size:.86rem;
    border:1px solid var(--line); border-radius:9px; background:#fff; color:var(--text);
  }
  .ed-f input[type=date]{flex:1 1 100%;}
  .ed-f input:focus{outline:0; border-color:var(--navy2);}
  .ed-t{display:flex; gap:6px;}
  .ed-b{
    border:1px solid var(--line); background:#fff; color:var(--muted);
    border-radius:999px; padding:3px 14px; font-size:.75rem; font-weight:600; cursor:pointer;
  }
  .ed-b.on{background:var(--navy); border-color:var(--navy); color:#fff;}
  .empty{
    text-align:center; color:var(--muted); font-size:.88rem; padding:20px 12px;
    border:1px dashed var(--line); border-radius:12px;
  }
  .actions{display:flex; gap:10px; margin-top:14px; flex-wrap:wrap;}
  .actions .btn{flex:1 1 190px;}

  /* ---------- tables ---------- */
  .tw{overflow-x:auto; -webkit-overflow-scrolling:touch;}
  table{width:100%; border-collapse:collapse; font-size:.92rem; min-width:430px;}
  thead th{
    background:var(--navy); color:#fff; font-size:.83rem; font-weight:600;
    padding:11px 12px; text-align:right; white-space:nowrap;
  }
  tbody td{padding:11px 12px; border-bottom:1px solid var(--line); vertical-align:middle;}
  tbody tr:last-child td{border-bottom:0;}
  tbody tr:hover td{background:#F8FAFD;}
  td.num{text-align:center; font-weight:700; font-variant-numeric:tabular-nums;
         white-space:nowrap; color:var(--navy);}
  .cell-cat{display:flex; align-items:center; gap:9px; font-weight:600;}
  .cell-cat .ic{width:17px; height:17px; color:var(--muted);}
  .share{display:flex; align-items:center; gap:8px; justify-content:center;}
  .share-bar{width:56px; height:6px; background:#E6ECF3; border-radius:99px; overflow:hidden;}
  .share-bar i{display:block; height:100%; background:var(--blue); border-radius:99px;}

  /* ---------- downloads ---------- */
  .dl{display:grid; grid-template-columns:repeat(auto-fit,minmax(250px,1fr)); gap:10px;}
  .dl a{
    display:flex; align-items:center; gap:11px; padding:13px 14px;
    border:1px solid var(--line); border-radius:12px; background:#fff;
    color:var(--text); text-decoration:none; font-size:.9rem; font-weight:600;
    transition:border-color .15s ease, box-shadow .15s ease;
  }
  .dl a:hover{border-color:var(--navy); box-shadow:var(--shadow); text-decoration:none;}
  .dl a .ic{color:var(--navy);}
  .dl a .sz{margin-inline-start:auto; font-size:.74rem; color:var(--muted); font-weight:500;}

  /* ---------- notes / footer ---------- */
  .note{
    display:flex; gap:10px; align-items:flex-start; margin-top:14px;
    background:var(--blue-bg); border:1px solid #D6E2F5; color:#28405F;
    border-radius:12px; padding:12px 14px; font-size:.86rem;
  }
  .note .ic{color:var(--blue); margin-top:3px;}
  .note.warn{background:var(--amber-bg); border-color:#EFDFAE; color:#6B5310;}
  .note.warn .ic{color:var(--amber);}
  footer{
    display:flex; align-items:center; justify-content:center; gap:8px;
    flex-wrap:wrap; text-align:center; color:var(--muted);
    font-size:.79rem; padding:20px 16px 26px;
  }
  footer .ic{width:14px; height:14px;}

  /* ---------- mobile ---------- */
  @media (max-width:640px){
    .hd{padding:15px 14px;}
    .brand h1{font-size:1.08rem;}
    .chip{display:none;}
    .wrap{padding:14px 12px 4px;}
    section{padding:14px; border-radius:12px;}
    .cards{grid-template-columns:repeat(2,1fr); gap:10px;}
    .card{padding:13px;}
    .card .v{font-size:1.22rem;}
    .ci{width:32px; height:32px; margin-bottom:8px;}
    .stats{grid-template-columns:repeat(2,1fr);}
    .mk{grid-template-columns:repeat(2,1fr); gap:8px;}
    .mk-i{padding:10px 11px;}
    .mk-v{font-size:.92rem;}
    .mk-cols{grid-template-columns:1fr;}
    #lg_d{width:100%;}
    #lg_a{width:calc(50% - 5px);}
    #lg_n{flex:1 1 calc(50% - 5px);}
    .form .btn{width:100%;}
    table{font-size:.86rem; min-width:360px;}
    thead th, tbody td{padding:9px 8px;}
    .share-bar{display:none;}
    .actions .btn{flex:1 1 100%;}
    .lock-card{padding:28px 20px;}
  }
</style>
</head>
<body class="locked">
$sprite

<div id="lock">
  <div class="lock-card">
    <div class="lock-logo">$i-lock</div>
    <h1>الخزنة المالية</h1>
    <p>أدخل رمز الدخول لعرض لوحتك</p>
    <input id="pw" type="password" inputmode="text" placeholder="••••" autocomplete="off" aria-label="رمز الدخول">
    <button class="btn lock-btn" onclick="go()">$i-log-in دخول</button>
    <div class="lock-err" id="err"></div>
    <div class="lock-hint">$i-info التحقق يتم داخل متصفحك — لا يُرسل الرمز لأي خادم</div>
  </div>
</div>

<header>
  <div class="hd">
    <div class="brand">
      <span class="logo">$i-wallet</span>
      <div>
        <h1>$name</h1>
        <p class="sub">لوحة الثروة · $p1 — $p2</p>
      </div>
    </div>
    <div class="chip">$i-shield-check تحديث تلقائي من رسائل البنك</div>
  </div>
</header>

<div class="wrap">

  <div class="cards">
    <div class="card">
      <span class="ci tone-red">$i-trending-down</span>
      <div class="v num">$net $cur</div>
      <div class="l">صافي المصروف (الفترة كاملة)</div>
    </div>
    <div class="card">
      <span class="ci tone-green">$i-trending-up</span>
      <div class="v num">$inc $cur</div>
      <div class="l">إجمالي الوارد</div>
    </div>
    <div class="card">
      <span class="ci tone-navy">$i-gauge</span>
      <div class="v num">$mo $cur</div>
      <div class="l">متوسط الصرف الشهري</div>
    </div>
    <div class="card">
      <span class="ci tone-gold">$i-piggy-bank</span>
      <div class="v num">$sav $cur</div>
      <div class="l">هدف ادخار $month</div>
    </div>
  </div>

  $market

  <section id="log">
    <div class="sec-h">$i-square-pen
      <h2>تسجيل سريع — سحب أو دخل</h2>
      <span class="hint">يُحفظ في هذا الجهاز</span>
    </div>

    <div class="seg" role="tablist">
      <button type="button" class="seg-btn on" id="seg_out" onclick="setType('out')">$i-trending-down سحب / مصروف</button>
      <button type="button" class="seg-btn t-in" id="seg_in" onclick="setType('in')">$i-trending-up دخل / وارد</button>
    </div>

    <div class="form">
      <input type="date" id="lg_d" aria-label="التاريخ">
      <input type="number" id="lg_a" placeholder="المبلغ ($cur)" inputmode="numeric" min="1" aria-label="المبلغ">
      <input type="text" id="lg_n" placeholder="ملاحظة (بنزين / راتب...)" maxlength="40" aria-label="ملاحظة">
      <button class="btn btn-pri" onclick="addEntry()">$i-plus إضافة</button>
    </div>

    <div class="stats">
      <div class="stat">
        <div class="sl">$i-gauge مصروف الأسبوع (من السبت)</div>
        <div class="sv num" id="st_wk">—</div>
        <div class="bar"><i id="bar_wk"></i></div>
      </div>
      <div class="stat">
        <div class="sl">$i-trending-up دخل الشهر</div>
        <div class="sv num v-green" id="st_in">—</div>
      </div>
      <div class="stat">
        <div class="sl">$i-trending-down مصروف الشهر</div>
        <div class="sv num v-red" id="st_out">—</div>
      </div>
      <div class="stat">
        <div class="sl">$i-wallet صافي الشهر</div>
        <div class="sv num" id="st_net">—</div>
      </div>
    </div>

    <div class="log-list" id="lg_list"></div>

    <div class="actions">
      <button class="btn btn-pri" onclick="copyReport()">$i-copy نسخ التقرير للمحاسبة</button>
      <button class="btn btn-danger" onclick="wipeLog()">$i-x مسح السجل</button>
    </div>

    <div class="note">$i-info التسجيلات محفوظة داخل متصفح هذا الجهاز فقط — لو فتحت من جهاز
    آخر هتلاقيه فاضي. أول ما تيجي للمحاسب: اضغط «نسخ التقرير» والصقه هنا بضغطة واحدة.</div>
  </section>

  <section>
    <div class="sec-h">$i-clipboard-list
      <h2>خطة $month — كل خانة بالجنيه</h2>
    </div>
    <div class="tw">
      <table>
        <thead><tr><th>البند</th><th class="num">شهريًا</th><th class="num">أسبوعيًا</th></tr></thead>
        <tbody>$plan_rows</tbody>
      </table>
    </div>
    <div class="note warn">$i-target سقف السحب النقدي الأسبوعي: <b class="num">$cap $cur</b> —
    سجّل كل سحبة في نفس اليوم من قسم «تسجيل سريع».</div>
  </section>

  <section>
    <div class="sec-h">$i-list-checks
      <h2>وين راحت الفلوس فعلاً — أكبر التصنيفات</h2>
    </div>
    <div class="tw">
      <table>
        <thead><tr><th>التصنيف</th><th class="num">الإجمالي ($cur)</th><th class="num">الحصة</th><th class="num">المعاملات</th></tr></thead>
        <tbody>$cat_rows</tbody>
      </table>
    </div>
  </section>

  <section>
    <div class="sec-h">$i-download
      <h2>تحميل تقاريري</h2>
    </div>
    <div class="dl">$dl</div>
  </section>

</div>
<footer>$i-shield-check بياناتك مبنية على جهازك وتُحدَّث من الرسايل الخام · تاريخ البناء $today</footer>

<script>
/* ===== gate: sha256 of the passcode — local check, nothing sent anywhere ===== */
var H = "$gate";
function sha(s) {
  var c = new TextEncoder().encode(s);
  return crypto.subtle.digest("SHA-256", c).then(function (b) {
    return Array.from(new Uint8Array(b)).map(function (x) {
      return x.toString(16).padStart(2, "0");
    }).join("");
  });
}
function unlock() {
  document.body.classList.remove("locked");
  document.getElementById("lock").style.display = "none";
}
function go() {
  var v = document.getElementById("pw").value;
  sha(v).then(function (d) {
    if (d === H) {
      unlock();
      try { sessionStorage.setItem("mf_ok", "1"); } catch (e) {}
    } else {
      document.getElementById("err").textContent = "الرمز غلط — جرّب تاني";
      document.getElementById("pw").value = "";
    }
  });
}
document.getElementById("pw").addEventListener("keydown", function (e) {
  if (e.key === "Enter") go();
});
try {
  if (sessionStorage.getItem("mf_ok") === "1") unlock();
} catch (e) {}

/* ===== quick log: expenses AND income, saved in localStorage ===== */
var CAP = $capnum;
var LKEY = "mf_log_$month";
var TYPE = "out";
var EDIT = -1, EDTYPE = "out";

function setType(t) {
  TYPE = t;
  document.getElementById("seg_out").classList.toggle("on", t === "out");
  document.getElementById("seg_in").classList.toggle("on", t === "in");
}
function localISO() {
  var d = new Date();
  return d.getFullYear() + "-" + String(d.getMonth() + 1).padStart(2, "0") +
         "-" + String(d.getDate()).padStart(2, "0");
}
function loadLog() {
  try {
    var a = JSON.parse(localStorage.getItem(LKEY) || "[]");
    return Array.isArray(a) ? a : [];
  } catch (e) { return []; }
}
function saveLog(a) {
  try { localStorage.setItem(LKEY, JSON.stringify(a)); } catch (e) {}
}
function fmtDate(iso) {
  var p = String(iso).split("-");
  return p[2] + "/" + p[1];
}
function weekStartISO() {
  var d = new Date();
  d.setHours(0, 0, 0, 0);
  d.setDate(d.getDate() - ((d.getDay() + 1) % 7));
  return d.getFullYear() + "-" + String(d.getMonth() + 1).padStart(2, "0") +
         "-" + String(d.getDate()).padStart(2, "0");
}
function escj(s) {
  return String(s).replace(/&/g, "&amp;").replace(/</g, "&lt;").replace(/>/g, "&gt;");
}
function escAttr(s) {
  return escj(s).replace(/"/g, "&quot;");
}
function totals() {
  var arr = loadLog(), ws = weekStartISO(), mpre = ws.slice(0, 7);
  var wk = 0, mIn = 0, mOut = 0;
  for (var i = 0; i < arr.length; i++) {
    var e = arr[i], isIn = (e.t === "in");
    if (e.d >= ws && !isIn) wk += e.a;
    if (String(e.d).slice(0, 7) === mpre) {
      if (isIn) mIn += e.a; else mOut += e.a;
    }
  }
  return { wk: wk, mIn: mIn, mOut: mOut, mpre: mpre, n: arr.length, arr: arr };
}
function renderLog() {
  var T = totals();

  document.getElementById("st_wk").textContent = Math.round(T.wk) + " / " + Math.round(CAP);
  var pct = CAP > 0 ? (T.wk / CAP) * 100 : 0;
  var bar = document.getElementById("bar_wk");
  bar.style.width = Math.min(100, pct) + "%";
  bar.className = pct > 100 ? "over" : (pct > 75 ? "warn" : "");

  document.getElementById("st_in").textContent = Math.round(T.mIn) + " $cur";
  document.getElementById("st_out").textContent = Math.round(T.mOut) + " $cur";
  var net = Math.round(T.mIn - T.mOut);
  var netEl = document.getElementById("st_net");
  netEl.textContent = (net > 0 ? "+" : "") + net + " $cur";
  netEl.className = "sv num " + (net >= 0 ? "v-green" : "v-red");

  var box = document.getElementById("lg_list");
  if (!T.n) {
    box.innerHTML = '<div class="empty">لا توجد تسجيلات بعد — أضف أول عملية من الأعلى</div>';
    return;
  }
  var rows = "";
  for (var i = T.arr.length - 1; i >= 0; i--) {
    var e = T.arr[i], isIn = (e.t === "in");
    if (i === EDIT) {
      rows += '<div class="log-row editing">' +
        '<span class="lic ' + (EDTYPE === "in" ? "tone-green" : "tone-red") + '">' +
          '<svg class="ic" aria-hidden="true"><use href="#ic-' +
          (EDTYPE === "in" ? "trending-up" : "trending-down") + '"/></svg></span>' +
        '<div class="li-main">' +
          '<div class="ed-f">' +
            '<input type="date" id="ed_d" value="' + String(e.d) +
              '" aria-label="التاريخ" onkeydown="if(event.key==\'Enter\')saveEdit();if(event.key==\'Escape\')cancelEdit();">' +
            '<input type="number" id="ed_a" value="' + e.a +
              '" min="1" inputmode="numeric" aria-label="المبلغ" onkeydown="if(event.key==\'Enter\')saveEdit();if(event.key==\'Escape\')cancelEdit();">' +
            '<input type="text" id="ed_n" value="' + escAttr(e.n || "") +
              '" maxlength="40" placeholder="ملاحظة" aria-label="ملاحظة" onkeydown="if(event.key==\'Enter\')saveEdit();if(event.key==\'Escape\')cancelEdit();">' +
          '</div>' +
          '<div class="ed-t">' +
            '<button type="button" class="ed-b' + (EDTYPE === "out" ? " on" : "") +
              '" id="edt_out" onclick="setEdType(\'out\')">سحب / مصروف</button>' +
            '<button type="button" class="ed-b' + (EDTYPE === "in" ? " on" : "") +
              '" id="edt_in" onclick="setEdType(\'in\')">دخل / وارد</button>' +
          '</div>' +
        '</div>' +
        '<button class="del ok" onclick="saveEdit()" aria-label="حفظ التعديل" title="حفظ">' +
          '<svg class="ic" aria-hidden="true"><use href="#ic-circle-check"/></svg></button>' +
        '<button class="del" onclick="cancelEdit()" aria-label="إلغاء" title="إلغاء">' +
          '<svg class="ic" aria-hidden="true"><use href="#ic-x"/></svg></button>' +
        '</div>';
      continue;
    }
    rows += '<div class="log-row">' +
      '<span class="lic ' + (isIn ? "tone-green" : "tone-red") + '">' +
        '<svg class="ic" aria-hidden="true"><use href="#ic-' +
        (isIn ? "trending-up" : "trending-down") + '"/></svg></span>' +
      '<div class="li-main">' +
        '<div class="li-t">' + (isIn ? "دخل" : "سحب") + " · " + fmtDate(e.d) + '</div>' +
        (e.n ? '<div class="li-s">' + escj(e.n) + '</div>' : '') +
      '</div>' +
      '<span class="li-amt ' + (isIn ? "v-green" : "v-red") + '">' +
        (isIn ? "+" : "-") + e.a + ' $cur</span>' +
      '<button class="del ed" onclick="editEntry(' + i + ')" aria-label="تعديل" title="تعديل">' +
        '<svg class="ic" aria-hidden="true"><use href="#ic-square-pen"/></svg></button>' +
      '<button class="del" onclick="delEntry(' + i + ')" aria-label="حذف" title="حذف">' +
        '<svg class="ic" aria-hidden="true"><use href="#ic-x"/></svg></button>' +
      '</div>';
  }
  box.innerHTML = rows;
}
function addEntry() {
  var a = parseInt(document.getElementById("lg_a").value, 10);
  var d = document.getElementById("lg_d").value || localISO();
  var n = document.getElementById("lg_n").value.trim();
  if (!a || a <= 0) {
    var f = document.getElementById("lg_a");
    f.style.borderColor = "var(--red)";
    f.focus();
    return;
  }
  document.getElementById("lg_a").style.borderColor = "";
  var arr = loadLog();
  arr.push({ d: d, a: a, n: n, t: TYPE });
  arr.sort(function (x, y) { return x.d < y.d ? -1 : (x.d > y.d ? 1 : 0); });
  saveLog(arr);
  EDIT = -1;
  document.getElementById("lg_a").value = "";
  document.getElementById("lg_n").value = "";
  document.getElementById("lg_a").focus();
  renderLog();
}
function delEntry(i) {
  var arr = loadLog();
  arr.splice(i, 1);
  saveLog(arr);
  if (EDIT === i) EDIT = -1;
  renderLog();
}
function editEntry(i) {
  var arr = loadLog();
  if (i < 0 || i >= arr.length) return;
  EDIT = i;
  EDTYPE = arr[i].t === "in" ? "in" : "out";
  renderLog();
  var f = document.getElementById("ed_a");
  if (f) { f.focus(); f.select(); }
}
function setEdType(t) {
  EDTYPE = t;
  var o = document.getElementById("edt_out"), n = document.getElementById("edt_in");
  if (o) o.classList.toggle("on", t === "out");
  if (n) n.classList.toggle("on", t === "in");
}
function cancelEdit() {
  EDIT = -1;
  renderLog();
}
function saveEdit() {
  var arr = loadLog();
  if (EDIT < 0 || EDIT >= arr.length) { EDIT = -1; renderLog(); return; }
  var f = document.getElementById("ed_a");
  var a = parseInt(f ? f.value : "", 10);
  if (!a || a <= 0) {
    if (f) { f.style.borderColor = "var(--red)"; f.focus(); }
    return;
  }
  if (f) f.style.borderColor = "";
  var dEl = document.getElementById("ed_d"), nEl = document.getElementById("ed_n");
  var d = (dEl && dEl.value) || arr[EDIT].d;
  var n = nEl ? nEl.value.trim() : (arr[EDIT].n || "");
  arr[EDIT] = { d: d, a: a, n: n, t: EDTYPE };
  arr.sort(function (x, y) { return x.d < y.d ? -1 : (x.d > y.d ? 1 : 0); });
  saveLog(arr);
  EDIT = -1;
  renderLog();
}
function wipeLog() {
  if (confirm("هتمسح كل تسجيلات الشهر ده — متأكد؟")) {
    saveLog([]);
    EDIT = -1;
    renderLog();
  }
}
function copyReport() {
  var T = totals();
  if (!T.n) { alert("لا توجد تسجيلات بعد"); return; }
  var outs = [], ins = [];
  for (var i = 0; i < T.arr.length; i++) {
    var e = T.arr[i], line = "- " + fmtDate(e.d) + ": " + e.a + " $cur" +
      (e.n ? " (" + e.n + ")" : "");
    if (e.t === "in") ins.push(line); else outs.push(line);
  }
  var txt = "تقرير التسجيل — الشهر " + T.mpre + "\n\n";
  if (ins.length) txt += "الدخل:\n" + ins.join("\n") + "\n\n";
  if (outs.length) txt += "المصروفات:\n" + outs.join("\n") + "\n\n";
  txt += "مصروف الأسبوع: " + Math.round(T.wk) + " من " + Math.round(CAP) +
         " (المتبقي " + Math.round(CAP - T.wk) + ")\n" +
         "الشهر: دخل " + Math.round(T.mIn) + " | مصروف " + Math.round(T.mOut) +
         " | الصافي " + (T.mIn - T.mOut > 0 ? "+" : "") +
         Math.round(T.mIn - T.mOut) + " $cur";
  function manualCopy(t) {
    var ta = document.createElement("textarea");
    ta.value = t;
    ta.setAttribute("readonly", "readonly");
    ta.style.cssText = "position:fixed;left:0;right:0;bottom:0;height:38vh;width:100%;" +
      "z-index:70;direction:rtl;text-align:right;font:14px/1.6 'Segoe UI',Tahoma,sans-serif;" +
      "padding:12px;border:0;border-top:3px solid #C9A227;background:#fff;" +
      "box-shadow:0 -8px 24px rgba(14,42,71,.25);";
    document.body.appendChild(ta);
    ta.focus(); ta.select();
    try { ta.setSelectionRange(0, t.length); } catch (e) {}
    var copied = false;
    try { copied = document.execCommand("copy"); } catch (e) {}
    if (copied) {
      document.body.removeChild(ta);
      alert("تم نسخ التقرير — افتح المحادثة والصقه");
    } else {
      ta.addEventListener("blur", function () {
        if (ta.parentNode) ta.parentNode.removeChild(ta);
      });
      alert("التقرير ظاهر في أسفل الشاشة — حدّده بالكامل وانسخه");
    }
  }
  var done = false;
  function finishCopy() { if (done) return; done = true; alert("تم نسخ التقرير — افتح المحادثة والصقه"); }
  function fallbackCopy() { if (done) return; done = true; manualCopy(txt); }
  if (navigator.clipboard && navigator.clipboard.writeText) {
    try {
      navigator.clipboard.writeText(txt).then(finishCopy, fallbackCopy);
      setTimeout(fallbackCopy, 1500);   // clipboard APIs may hang without user gesture
    } catch (e) { fallbackCopy(); }
  } else {
    manualCopy(txt);
  }
}
(function initLog() {
  var di = document.getElementById("lg_d");
  if (di && !di.value) di.value = localISO();
  renderLog();
})();
</script>
</body>
</html>
"""


def main():
    s = json.load(open(SUMMARY, encoding="utf-8"))
    os.makedirs(WEB, exist_ok=True)

    # ---- plan rows from xlsx ----
    plan_rows = []
    savings = 0.0
    weekly_cap = 0.0
    if os.path.isfile(PLAN):
        import openpyxl
        wb = openpyxl.load_workbook(PLAN, data_only=True)
        ws = wb[wb.sheetnames[0]]
        for row in ws.iter_rows(min_row=4, max_row=30, max_col=5):
            label, amt = row[1].value, row[3].value
            if not label or not isinstance(amt, (int, float)):
                continue
            if label in ("الإجمالي", "تفاوت التقريب (لو فيه)"):
                continue
            plan_rows.append((label, amt))
            if "ادخار" in str(label) or "فائض" in str(label):
                savings += float(amt)
            if "ATM" in str(label) or "سحب" in str(label):
                weekly_cap = float(amt) / 4.33

    # ---- categories (amount, count) + share of net spend ----
    cats = s.get("categories") or {}
    cat_rows = sorted(cats.items(), key=lambda kv: kv[1]["amount"], reverse=True)
    total_out = float(s.get("net_out") or 0) or 1.0

    # ---- plan table html ----
    plan_html = "\n".join(
        '<tr><td><span class="cell-cat">{ic}{lbl}</span></td>'
        '<td class="num">{amt}</td><td class="num">{wk}</td></tr>'.format(
            ic=svg(icon_for(lbl)), lbl=esc(lbl),
            amt=money(amt), wk=money(float(amt) / 4.33))
        for lbl, amt in plan_rows)

    cat_html = "\n".join(
        '<tr><td><span class="cell-cat">{ic}{name}</span></td>'
        '<td class="num">{amt}</td>'
        '<td class="num"><span class="share">'
        '<span class="share-bar"><i style="width:{pct:.0f}%"></i></span>{pct:.1f}%</span></td>'
        '<td class="num">{n}</td></tr>'.format(
            ic=svg(icon_for(name)), name=esc(name),
            amt=money(v["amount"]), pct=v["amount"] / total_out * 100, n=money(v["n"]))
        for name, v in cat_rows)

    # ---- market pulse section (from 02_Reports/market_watch.json) ----
    market_html = ""
    mkt = os.path.join(BASE, "02_Reports", "market_watch.json")
    if os.path.isfile(mkt):
        try:
            m = json.load(open(mkt, encoding="utf-8"))
            chips = "\n".join(
                '<div class="mk-i"><span class="mk-l">{ic}{lab}</span>'
                '<b class="mk-v num">{val}</b><span class="mk-s">{src}</span></div>'.format(
                    ic=svg(it.get("icon") or "circle-dot"), lab=esc(it.get("label", "")),
                    val=esc(it.get("value", "")), src=esc(it.get("src", "")))
                for it in m.get("items", []))
            impact = "\n".join(
                "<li>{ic}<span>{t}</span></li>".format(ic=svg("target"), t=esc(t))
                for t in m.get("impact", []))
            news = "\n".join(
                '<li>{ic}<span>{t}<span class="mk-date">{d} — {w}</span></span></li>'.format(
                    ic=svg("circle-alert"), t=esc(n.get("title", "")),
                    d=esc(n.get("date", "")), w=esc(n.get("why", "")))
                for n in m.get("news", []))
            if chips:
                market_html = (
                    '<section id="market"><div class="sec-h">{ic}'
                    '<h2>نبض السوق — {city}</h2>'
                    '<span class="hint">آخر فحص: {upd} — تحديث أسبوعي</span></div>'
                    '<div class="mk">{chips}</div>'
                    '<div class="mk-cols"><div><div class="mk-h">{ti} أثر الأسعار على خطتك</div>'
                    '<ul class="mk-ul">{impact}</ul></div>'
                    '<div><div class="mk-h">{ni} أخبار تؤثر في مصروفاتك</div>'
                    '<ul class="mk-ul">{news}</ul></div></div>'
                    '<div class="note">{ii} الأرقام دي بتتحدّث أسبوعيًا كل سبت وأمام أي خطة شهرية، '
                    'والمصدر الكامل بتواريخه في وثيقة الأسعار داخل المجلد المحلي.</div></section>'
                ).format(ic=svg("refresh-cw"), upd=esc(m.get("last_checked", "")),
                         chips=chips, impact=impact, news=news, city=esc(CFG.get("city") or ""),
                         ti=svg("target"), ni=svg("circle-alert"), ii=svg("info"))
        except Exception as e:
            print("WARN market section skipped: {}".format(e))

    # ---- download links (token-prefixed names, with size) ----
    dl_items = []
    for f, orig, t in DL:
        src = os.path.join(BASE, "02_Reports", f)
        if os.path.isfile(src):
            ext = os.path.splitext(orig)[1].lower()
            ic = {".xlsx": "file-spreadsheet", ".csv": "table",
                  ".json": "braces", ".html": "layout-dashboard"}.get(ext, "file-text")
            if orig.startswith("plan_"):
                ic = "calendar-check"
            dl_items.append((token_for(f) + "_" + f, t, ic,
                             human(os.path.getsize(src))))
    dl_html = "\n".join(
        '<a href="{s}" download>{ic}{t}<span class="sz">{sz}</span></a>'.format(
            s=esc(sn), ic=svg(ic), t=esc(t), sz=sz)
        for sn, t, ic, sz in dl_items)

    period = str(s.get("period") or "")
    if ".." in period:
        p_start, p_end = [x.strip() for x in period.split("..", 1)]
    else:
        p_start = p_end = period

    html = Template(expand_icons(TEMPLATE)).safe_substitute(
        sprite=sprite(),
        p1=esc(p_start), p2=esc(p_end),
        net=money(s.get("net_out")), inc=money(s.get("total_credits")),
        mo=money(s.get("avg_monthly_out")), sav=money(savings),
        month=esc(MONTH), cap=money(weekly_cap),
        capnum=round(weekly_cap, 2),
        name=esc(CFG.get("name") or "MyFinance"),
        cur=esc(CFG.get("currency") or "EGP"),
        plan_rows=plan_html, cat_rows=cat_html, market=market_html,
        dl=dl_html, today=dt.date.today().strftime("%d/%m/%Y"),
        gate=GATE,
    )
    if "$i-" in html or "$market" in html:
        print("WARN: unresolved template token in output")

    with open(os.path.join(WEB, "index.html"), "w", encoding="utf-8") as f:
        f.write(html)

    # ---- copy downloadable reports (token-prefixed names) ----
    copied = 0
    for stored, _, _, _ in dl_items:
        src_name = stored.split("_", 1)[1] if "_" in stored else stored
        # match original by token map
        for f, orig, _t in DL:
            if token_for(f) + "_" + f == stored:
                src_name = f
                break
        src = os.path.join(BASE, "02_Reports", src_name)
        if os.path.isfile(src):
            shutil.copy2(src, os.path.join(WEB, stored))
            copied += 1

    print("WEB OK  index.html written  ({} plan rows, {} categories, {} icons)".format(
        len(plan_rows), len(cat_rows), len(ICONS)))
    print("copied {} report files into 06_Web".format(copied))
    print("open: " + os.path.join(WEB, "index.html"))


if __name__ == "__main__":
    main()
